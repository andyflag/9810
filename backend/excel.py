"""Original workbook import and compatible, fully calculated Excel export."""

from datetime import date, datetime
from io import BytesIO
import json, zipfile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .domain import SCHEMA, ORDER, RuleError, clean, DEFAULT_SETTINGS, compute

MAX_XLSX = 10 * 1024 * 1024


def parse(content):
    if len(content) > MAX_XLSX:
        raise RuleError("Excel 文件不得超过 10 MB")
    try:
        with zipfile.ZipFile(BytesIO(content)) as z:
            if (
                sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024
                or len(z.infolist()) > 1000
            ):
                raise RuleError("Excel 解压后过大")
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except RuleError:
        raise
    except Exception:
        raise RuleError("无法读取文件，请上传有效的 .xlsx 工作簿") from None
    result = {k: [] for k in ORDER}
    errors, recognized, count = [], 0, 0
    settings = {}
    try:
        for kind in ORDER:
            spec = SCHEMA[kind]
            if spec["sheet"] not in wb.sheetnames:
                continue
            recognized += 1
            sheet = wb[spec["sheet"]]
            if sheet.max_row and sheet.max_row > 20000:
                raise RuleError("单个工作表最多 20,000 行")
            if sheet.max_column and sheet.max_column > 100:
                raise RuleError("工作表列数超出模板范围")
            header = spec["header_row"]
            headers = next(sheet.iter_rows(min_row=header, max_row=header, values_only=True), ())
            normalize = lambda s: str(s or "").replace("（自动）", "").strip()
            # Ignore new optional columns absent from the original workbook.
            bad = [
                f["label"]
                for f in spec["fields"]
                if f["required"]
                and (
                    len(headers) < f["column"]
                    or normalize(headers[f["column"] - 1]) != normalize(f["label"])
                )
            ]
            if bad:
                raise RuleError(f"{spec['sheet']} 表头不匹配：{'、'.join(bad)}")
            for rowno, row in enumerate(
                sheet.iter_rows(
                    min_row=header + 1, max_col=max(f["column"] for f in spec["fields"])
                ),
                header + 1,
            ):
                payload = {}
                for f in spec["fields"]:
                    cell = row[f["column"] - 1]
                    if f["computed"] or cell.value is None:
                        continue
                    # Original blank rows carry formulas; never evaluate uploaded formulas.
                    if cell.data_type == "f":
                        raise RuleError(f"{spec['sheet']} 第 {rowno} 行：手填列不能包含公式")
                    value = cell.value
                    if isinstance(value, (datetime, date)):
                        value = value.strftime("%Y-%m-%d")
                    if f["type"] == "text":
                        if (
                            isinstance(value, (int, float))
                            and f["key"] in ("declaration_no", "sku", "invoice_no", "settlement_id")
                            and len(str(int(value))) > 15
                        ):
                            raise RuleError(
                                f"{spec['sheet']} 第 {rowno} 行：长编号必须保存为文本，避免 Excel 精度丢失"
                            )
                        value = str(value)
                    payload[f["key"]] = value
                if not payload:
                    continue
                if kind == "skus" and not payload.get("sku") and not payload.get("name"):
                    continue
                try:
                    result[kind].append(clean(kind, payload))
                except RuleError as e:
                    errors.append(f"{spec['sheet']} 第 {rowno} 行：{e}")
                count += 1
                if count > 10000:
                    raise RuleError("每次导入最多 10,000 条业务记录")
        if "参数表" in wb.sheetnames:
            s = wb["参数表"]
            for k, pos in [
                ("company_name", "B5"),
                ("company_type", "B6"),
                ("shop_name", "B7"),
                ("tax_authority", "B8"),
                ("filing_date", "B9"),
            ]:
                v = s[pos].value
                if v is not None and s[pos].data_type != "f":
                    settings[k] = (
                        v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else str(v)
                    )
    finally:
        wb.close()
    if not recognized:
        raise RuleError("未找到原模板工作表，请下载模板后填写")
    if errors:
        raise RuleError("\n".join(errors[:20]))
    return result, settings


def fingerprint(kind, row):
    return json.dumps(clean(kind, row), ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def merge(current, incoming):
    combined = {k: list(current[k]) for k in ORDER}
    new, skipped = [], 0
    for kind in ORDER:
        seen = {fingerprint(kind, r) for r in current[kind]}
        for r in incoming[kind]:
            mark = fingerprint(kind, r)
            if mark in seen:
                skipped += 1
                continue
            seen.add(mark)
            combined[kind].append(r)
            new.append((kind, r))
    return combined, new, skipped


def safe_cell(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Formula injection protection, even for =,+,-,@.


def export(data, settings):
    result = compute(data, settings)
    wb = Workbook()
    wb.remove(wb.active)
    notes = wb.create_sheet("使用说明")
    notes.append(["9810 出口海外仓核销台账 · 应用导出"])
    notes.append(["导出包含当前核算结果；在应用中修改记录后重新导出可刷新数值。"])
    notes.append(["按原模板工作表、表头行和列顺序导出，可再次导入。公式列由服务端计算。"])
    notes.append(["跨币种分别统计；部分收汇按金额比例分配；陆续留存中不视为单证齐全。"])
    summary = wb.create_sheet("核销总表")
    summary.append(["核销总表"])
    summary.append(["自动汇总所有业务台账"])
    summary.append(
        ["报关项数", result["stats"]["export_count"], "报关数量", result["stats"]["quantity"]]
    )
    summary.append([])
    summary.append([])
    mapping = [
        ("核销编号", "ledger_id"),
        ("SKU", "sku"),
        ("品名", "name"),
        ("出口日期", "export_date"),
        ("报关数量", "quantity"),
        ("签收数量", "received_qty"),
        ("入仓差异", "warehouse_difference"),
        ("累计净售出", "net_sold"),
        ("异常减少", "exception_qty"),
        ("在仓未售", "remaining"),
        ("销售进度（%）", "progress"),
        ("累计净回款（原币）", "net_revenue"),
        ("其中已收汇（原币）", "received_amount"),
        ("收汇截止日", "receipt_deadline"),
        ("收汇状态", "receipt_status"),
        ("已申报退税额（元）", "declared_tax"),
        ("其中预退税额（元）", "prepaid_tax"),
        ("核算期末", "tax_deadline"),
        ("核算状态", "tax_status"),
        ("单证情况", "document_status"),
        ("需要关注", "attention"),
        ("币种", "currency"),
    ]
    summary.append([a for a, b in mapping])
    for i, r in enumerate(result["ledger"], 7):
        r["attention"] = "、".join(a["text"] for a in r["alerts"]) or "正常"
        for j, (_, key) in enumerate(mapping, 1):
            safe_cell(summary.cell(i, j), r.get(key))
    for kind in ORDER:
        spec = SCHEMA[kind]
        s = wb.create_sheet(spec["sheet"])
        s.cell(1, 1, spec["label"])
        s.cell(2, 1, spec["description"])
        if kind == "skus":
            for key, label, row in [
                ("company_name", "企业名称", 5),
                ("company_type", "企业类型", 6),
                ("shop_name", "亚马逊店铺名称", 7),
                ("tax_authority", "主管税务机关", 8),
                ("filing_date", "出口退（免）税备案日期", 9),
            ]:
                s.cell(row, 1, label)
                safe_cell(s.cell(row, 2), settings.get(key, ""))
        for f in spec["fields"]:
            s.cell(
                spec["header_row"], f["column"], f["label"] + ("（自动）" if f["computed"] else "")
            )
        for i, r in enumerate(result["records"][kind], spec["header_row"] + 1):
            for f in spec["fields"]:
                value = r.get(f["key"])
                cell = s.cell(i, f["column"])
                if value is not None and f["type"] in ("number", "integer"):
                    value = float(value)
                if value and f["type"] == "date":
                    value = date.fromisoformat(value)
                safe_cell(cell, value)
                if f["type"] == "date":
                    cell.number_format = "yyyy-mm-dd"
                if f["key"] == "tax_rate":
                    cell.number_format = "0.00%"
                if f["computed"]:
                    cell.fill = PatternFill("solid", fgColor="ECF5F5")
    for s in wb:
        h = 6 if s.title == "核销总表" else 12 if s.title == "参数表" else 3
        if s.title == "使用说明":
            s.column_dimensions["A"].width = 100
            continue
        s.freeze_panes = f"C{h + 1}"
        s.auto_filter.ref = f"A{h}:{s.cell(max(s.max_row, h), s.max_column).coordinate}"
        for c in s[h]:
            c.fill = PatternFill("solid", fgColor="136D58")
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            s.column_dimensions[c.column_letter].width = 22 if c.column > 1 else 30
        s.row_dimensions[h].height = 36
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
