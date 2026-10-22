from datetime import date
from io import BytesIO
from pathlib import Path
import json
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend import db, excel
from backend.main import app
from backend.domain import (
    ROOT,
    ORDER,
    DEFAULT_SETTINGS,
    clean,
    compute,
    validate_state,
    RuleError,
)

HEADERS = {"X-Requested-With": "9810"}


class ScopedClient(TestClient):
    """Model the UI selecting the workspace returned by identity operations."""

    def request(self, method, url, *args, **kwargs):
        response = super().request(method, url, *args, **kwargs)
        path = response.request.url.path
        if response.status_code == 200 and path in (
            "/api/auth/register",
            "/api/auth/login",
            "/api/workspaces/switch",
            "/api/invitations/accept",
        ):
            self.headers["X-Workspace-ID"] = response.json()["tenant_id"]
        if response.status_code == 200 and path == "/api/auth/logout":
            self.headers.pop("X-Workspace-ID", None)
        return response


@pytest.fixture
def sample():
    source = json.loads((ROOT / "sample.json").read_text())
    return {k: [clean(k, r) for r in source[k]] for k in ORDER}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DATA_DIR", tmp_path)
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.sqlite3")
    with ScopedClient(app, headers=HEADERS) as c:
        assert (
            c.post(
                "/api/auth/register",
                json={
                    "email": "owner@example.com",
                    "password": "safe-password-9810",
                    "company_name": "测试企业",
                },
            ).status_code
            == 200
        )
        yield c


def state(c):
    r = c.get("/api/state")
    assert r.status_code == 200
    return r.json()


def test_original_workbook_values(sample):
    r = compute(sample, today=date(2026, 10, 22))
    assert r["stats"]["quantity"] == 1500
    assert r["stats"]["net_sold"] == 885
    assert r["stats"]["remaining"] == 608
    assert r["stats"]["prepaid_tax"] == 7540
    assert r["stats"]["currencies"]["USD"] == {
        "net_revenue": 17391,
        "received": 10195,
        "pending": 7196,
    }
    assert [x["remaining"] for x in r["ledger"]] == [308, 300]
    assert all(x["allocation_difference"] == 0 for x in r["records"]["receipts"])
    assert r["records"]["documents"][0]["completeness"] == "留存中（2项）"
    assert r["records"]["documents"][0]["timeliness"] == "已超期未完成"


def test_partial_receipts_do_not_become_fully_received(sample):
    sample["receipts"][0]["received_amount"] = "5097.50"
    r = compute(sample)
    assert r["stats"]["currencies"]["USD"]["received"] == 5097.5
    assert r["records"]["sales"][0]["receipt_status"] == "部分收汇"
    assert r["records"]["sales"][0]["received_allocated"] == 2698.5


def test_zero_stock_and_inventory_guard(sample):
    sample["exports"][0]["received_qty"] = 0
    assert compute(sample)["ledger"][0]["remaining"] == -688
    with pytest.raises(RuleError, match="库存不足"):
        validate_state(sample, DEFAULT_SETTINGS)


def test_decimal_rounding_and_company_type(sample):
    sample["exports"][0]["fob"] = "0.35"
    sample["exports"][0]["exchange_rate"] = "7.1"
    r = compute(sample, {**DEFAULT_SETTINGS, "company_type": "生产企业"})
    assert r["records"]["exports"][0]["fob_cny"] == 2.49
    assert r["records"]["exports"][0]["expected_tax"] == 0.32
    assert compute(sample)["records"]["exports"][0]["expected_tax"] == 4680


def test_deadline_edges_and_override(sample):
    row = sample["taxes"][0]
    assert compute(sample, today=date(2027, 3, 1))["records"]["taxes"][0]["status"] == "临近截止"
    assert compute(sample, today=date(2027, 4, 30))["records"]["taxes"][0]["status"] == "临近截止"
    assert compute(sample, today=date(2027, 5, 1))["records"]["taxes"][0]["status"] == "已逾期"
    row["deadline_override"] = "2027-04-15"
    assert compute(sample, today=date(2027, 4, 16))["records"]["taxes"][0]["status"] == "已逾期"
    row["reconciled_date"] = "2027-04-14"
    assert compute(sample, today=date(2027, 5, 1))["records"]["taxes"][0]["status"] == "已核算"


def test_multiple_prepayments_share_sales_once(sample):
    first = sample["taxes"][0]
    first["quantity"] = 500
    first["amount"] = "2340"
    second = {**first, "declaration_date": "2026-04-01", "reference": "YT-second"}
    sample["taxes"].insert(1, second)
    r = compute(sample)
    taxes = [t for t in r["records"]["taxes"] if t["ledger_id"] == first["ledger_id"]]
    assert [t["sold_qty"] for t in taxes] == [500, 185]
    assert sum(t["reference_amount"] for t in taxes) == 3205.8


def test_currency_isolation(sample):
    # An independent EUR lot has no effect on USD receipt totals.
    e = {**sample["exports"][1], "declaration_no": "EU-2026", "currency": "EUR"}
    sample["exports"].append(e)
    r = compute(sample)
    assert set(r["stats"]["currencies"]) == {"EUR", "USD"}
    sample["receipts"][0]["currency"] = "EUR"
    with pytest.raises(RuleError, match="币种"):
        validate_state(sample, DEFAULT_SETTINGS)


def test_excel_roundtrip_and_formula_injection(sample):
    sample["skus"][0]["notes"] = '=HYPERLINK("https://example.com","test")'
    content = excel.export(sample, DEFAULT_SETTINGS)
    wb = load_workbook(BytesIO(content), data_only=False)
    assert wb["参数表"]["I13"].data_type == "s"
    restored, _ = excel.parse(content)
    assert compute(restored)["stats"] == compute(sample)["stats"]
    merged, new, skipped = excel.merge(sample, restored)
    assert not new and skipped == 13


def test_invalid_workbook_and_required_fields(sample):
    with pytest.raises(RuleError):
        excel.parse(b"not an xlsx")
    with pytest.raises(RuleError):
        clean("exports", {"quantity": -1})
    with pytest.raises(RuleError):
        clean("skus", {"sku": "X", "name": "X", "unit": "件", "tax_rate": "NaN"})
    with pytest.raises(RuleError):
        clean("skus", {"sku": "X", "name": "X", "unit": "件", "tax_rate": "1.3"})


def test_cookie_auth_csrf_logout_and_validation(client):
    assert state(client)["user"]["role"] == "owner"
    assert client.post("/api/demo", headers={"X-Requested-With": ""}).status_code == 403
    assert client.post("/api/demo").status_code == 200
    assert client.post("/api/demo").status_code == 409
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/state").status_code == 401
    assert (
        client.post(
            "/api/auth/login", json={"email": "owner@example.com", "password": "wrong-password"}
        ).status_code
        == 401
    )
    r = client.post(
        "/api/auth/login", json={"email": "owner@example.com", "password": "safe-password-9810"}
    )
    assert r.status_code == 200
    assert "HttpOnly" in r.headers["set-cookie"] and "SameSite=lax" in r.headers["set-cookie"]


def test_tenant_isolation_for_records_and_files(client):
    client.post("/api/demo")
    a = state(client)
    row = a["records"]["skus"][0]
    doc = a["records"]["documents"][0]
    r = client.post(
        "/api/documents/" + doc["id"] + "/files",
        files={"file": ("contract.txt", b"company A confidential")},
    )
    assert r.status_code == 200
    file_id = r.json()["id"]
    client.post(
        "/api/auth/register", json={"email": "other@example.com", "password": "safe-password-9810"}
    )
    assert state(client)["stats"]["export_count"] == 0
    assert (
        client.put("/api/records/skus/" + row["id"], json={"data": row, "version": 1}).status_code
        == 404
    )
    assert client.delete("/api/records/skus/" + row["id"] + "?version=1").status_code == 404
    assert client.get("/api/files/" + file_id).status_code == 404
    assert (
        client.post(
            "/api/documents/" + doc["id"] + "/files", files={"file": ("bad.txt", b"bad")}
        ).status_code
        == 404
    )
    assert (
        client.post(
            "/api/workspaces/switch", json={"tenant_id": a["user"]["tenant_id"]}
        ).status_code
        == 403
    )


def test_viewer_invitation_and_session_revocation(client):
    invited = client.post(
        "/api/invitations", json={"email": "viewer@example.com", "role": "viewer"}
    ).json()["token"]
    owner_cookie = client.cookies.get("session9810")
    client.post(
        "/api/auth/register", json={"email": "viewer@example.com", "password": "safe-password-9810"}
    )
    viewer_id = state(client)["user"]["id"]
    assert client.post("/api/invitations/accept", json={"token": invited}).status_code == 200
    viewer_cookie = client.cookies.get("session9810")
    assert state(client)["user"]["role"] == "viewer"
    assert client.post("/api/demo").status_code == 403
    assert client.post("/api/invitations", json={"email": "a@example.com"}).status_code == 403
    assert client.post("/api/invitations/accept", json={"token": invited}).status_code == 422
    assert client.get("/api/export").status_code == 200
    client.cookies.clear()
    client.cookies.set("session9810", owner_cookie)
    assert client.delete("/api/team/" + viewer_id).status_code == 200
    client.cookies.clear()
    client.cookies.set("session9810", viewer_cookie)
    assert client.get("/api/state").status_code == 403


def test_crud_concurrency_and_dependency_guards(client):
    client.post("/api/demo")
    s = state(client)
    row = s["records"]["exports"][0]
    assert client.delete("/api/records/exports/" + row["id"] + "?version=1").status_code == 422
    row["notes"] = "已核对"
    assert (
        client.put(
            "/api/records/exports/" + row["id"], json={"data": row, "version": 1}
        ).status_code
        == 200
    )
    assert (
        client.put(
            "/api/records/exports/" + row["id"], json={"data": row, "version": 1}
        ).status_code
        == 409
    )
    sale = {**s["records"]["sales"][0], "sold_qty": 10000}
    assert client.post("/api/records/sales", json={"data": sale}).status_code == 422
    assert len(state(client)["records"]["sales"]) == 3
    assert any(a["action"] == "更新记录" for a in client.get("/api/audit").json())


def test_import_preview_atomicity_and_dedup(client):
    path = Path(__file__).parents[1] / "docs/亚马逊9810核销台账模板-初版.xlsx"
    content = path.read_bytes()
    r = client.post("/api/import?dry_run=true", files={"file": ("template.xlsx", content)})
    assert r.status_code == 200, r.text
    assert r.json()["total"] == 13
    assert not state(client)["records"]["exports"]
    assert (
        client.post(
            "/api/import?dry_run=false", files={"file": ("template.xlsx", content)}
        ).status_code
        == 200
    )
    r = client.post("/api/import?dry_run=false", files={"file": ("template.xlsx", content)})
    assert r.json()["skipped"] == 13 and r.json()["total"] == 0
    workbook = load_workbook(BytesIO(content))
    workbook["①出口报关"]["I4"] = "BAD-SKU"
    stream = BytesIO()
    workbook.save(stream)
    assert (
        client.post(
            "/api/import?dry_run=false", files={"file": ("bad.xlsx", stream.getvalue())}
        ).status_code
        == 422
    )
    assert state(client)["stats"]["quantity"] == 1500


def test_fifo_splits_lots_and_preserves_cents(client):
    client.post("/api/demo")
    s = state(client)
    second = {
        **next(r for r in s["records"]["exports"] if r["sku"] == "DEMO-A01"),
        "declaration_no": "FIFO-LOT-2",
        "export_date": "2026-04-01",
        "quantity": 100,
        "received_qty": 100,
    }
    assert client.post("/api/records/exports", json={"data": second}).status_code == 200
    assert (
        client.post(
            "/api/records/receipts",
            json={
                "data": {
                    "settlement_id": "FIFO-SETTLEMENT",
                    "currency": "USD",
                    "site": "美国US",
                    "gross": "100.01",
                }
            },
        ).status_code
        == 200
    )
    payload = {
        "sku": "DEMO-A01",
        "settlement_id": "FIFO-SETTLEMENT",
        "month": "2026-10-01",
        "quantity": 350,
        "gross": "100.01",
        "fees": "7.01",
    }
    r = client.post("/api/sales/fifo", json=payload)
    assert r.status_code == 200, r.text
    preview = r.json()
    assert [r["sold_qty"] for r in preview["rows"]] == [308, 42]
    assert sum(float(r["gross"]) for r in preview["rows"]) == 100.01
    assert (
        client.post(
            "/api/sales/fifo", json={**payload, "commit": True, "revision": preview["revision"] - 1}
        ).status_code
        == 409
    )
    assert (
        client.post(
            "/api/sales/fifo", json={**payload, "commit": True, "revision": preview["revision"]}
        ).status_code
        == 200
    )
    assert len(state(client)["records"]["sales"]) == 5
    assert client.post("/api/sales/fifo", json={**payload, "quantity": 9999}).status_code == 422


def test_settings_optimistic_lock_and_files_persistence(client):
    s = state(client)
    body = {"data": {**s["settings"], "company_type": "生产企业"}, "version": s["revision"]}
    assert client.put("/api/settings", json=body).status_code == 200
    assert client.put("/api/settings", json=body).status_code == 409
    client.post("/api/demo")
    s = state(client)
    doc = s["records"]["documents"][0]
    r = client.post(
        "/api/documents/" + doc["id"] + "/files", files={"file": ("invoice.txt", b"proof")}
    )
    file_id = r.json()["id"]
    assert client.get("/api/files/" + file_id).content == b"proof"
    db.initialize()  # Restart-safe schema initialization must preserve records.
    assert state(client)["stats"]["quantity"] == 1500
    assert client.get("/api/files/" + file_id).content == b"proof"
    assert client.delete("/api/files/" + file_id).status_code == 200
    assert client.get("/api/files/" + file_id).status_code == 404


def test_overallocation_cannot_invent_received_cash(sample):
    sample["sales"][0]["gross"] = "9997"
    result = compute(sample)
    assert result["records"]["receipts"][0]["allocation_difference"] == -1000
    assert result["stats"]["currencies"]["USD"]["received"] == 10195
    assert result["records"]["sales"][0]["receipt_status"] == "部分收汇"


def test_evidence_date_before_sale_is_rejected(sample):
    row = {**sample["sales"][0], "evidence_date": "2026-05-01"}
    with pytest.raises(RuleError, match="不得早于"):
        clean("sales", row)
