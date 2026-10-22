"""Financial and workbook regressions identified in the full application audit."""

from io import BytesIO
from pathlib import Path
import json
import pytest
from openpyxl import load_workbook
from backend.domain import (
    ORDER,
    SCHEMA,
    DEFAULT_SETTINGS,
    NORMAL,
    RuleError,
    allocate_money,
    clean,
    compute,
    dec,
    validate_state,
)
from backend import excel


def sample():
    source = json.loads((Path(__file__).parents[1] / "backend/sample.json").read_text())
    return {kind: [clean(kind, row) for row in source[kind]] for kind in ORDER}


def test_cent_allocation_conserves_cash_and_signed_refunds():
    assert allocate_money("0.02", [1, 1, 1, 1]) == [0.01, 0.01, 0, 0]
    for amount in ["0.02", "-0.02", "10.005", "-10.005", "0.005", "-0.005"]:
        result = allocate_money(amount, [1, 1, 1, 1])
        from decimal import ROUND_HALF_UP

        assert sum(map(dec, result)) == dec(amount).quantize(dec(".01"), rounding=ROUND_HALF_UP)
        assert all(value <= 0 if dec(amount) < 0 else value >= 0 for value in result)
    signed = allocate_money("0.02", ["0.04", "-0.01", "0.01"])
    assert signed[0] >= 0 and signed[1] <= 0 and signed[2] >= 0
    assert sum(map(dec, signed)) == dec(".02")
    assert allocate_money("10", [1, 1], 4) == [2.5, 2.5]


def test_four_one_cent_sales_do_not_generate_negative_cash():
    data = sample()
    data["sales"] = [
        dict(data["sales"][0], sold_qty=1, returned_qty=0, gross="0.01", fees="0") for _ in range(4)
    ]
    data["receipts"] = [
        dict(
            data["receipts"][0],
            gross="0.04",
            commission="0",
            fulfillment="0",
            storage="0",
            advertising="0",
            refunds="0",
            other_fees="0",
            received_amount="0.02",
        )
    ]
    result = validate_state(data, DEFAULT_SETTINGS)
    assert [row["received_allocated"] for row in result["records"]["sales"]] == [0.01, 0.01, 0, 0]


def test_same_day_tax_allocation_survives_excel_round_trip():
    data = sample()
    first = data["taxes"][0]
    data["taxes"] = [
        dict(first, id="z", reference="FIRST", quantity=500, amount="500"),
        dict(first, id="a", reference="SECOND", quantity=500, amount="1000"),
    ]

    def amounts(records):
        return [
            (row["reference"], row["sold_qty"], row["reference_amount"])
            for row in compute(records)["records"]["taxes"]
        ]

    assert amounts(data) == [("FIRST", 500, 500), ("SECOND", 185, 370)]
    restored, _ = excel.parse(excel.export(data, DEFAULT_SETTINGS))
    assert amounts(restored) == amounts(data)


@pytest.mark.parametrize(
    "normal,sold,expected",
    [
        ([-100], 685, 685),
        ([100, -100], 685, 685),
        ([100], 685, 585),
        ([1000], 685, 0),
        ([-100], 0, 0),
    ],
)
def test_normal_reversals_cannot_expand_prepaid_allocation(normal, sold, expected):
    data = sample()
    first = data["taxes"][0]
    data["taxes"] = [
        dict(first, reference="PRE-1", quantity=500, amount="2340"),
        dict(first, reference="PRE-2", quantity=500, amount="2340"),
        *[
            dict(
                first,
                method=NORMAL,
                reference=f"NORMAL-{i}",
                quantity=quantity,
                amount=format(dec(quantity) * dec("4.68"), "f"),
            )
            for i, quantity in enumerate(normal)
        ],
    ]
    if not sold:
        for row in data["sales"]:
            if row["ledger_id"] == first["ledger_id"]:
                row["returned_qty"] = row["sold_qty"]
    state = validate_state(data, DEFAULT_SETTINGS)
    prepaid = state["records"]["taxes"][:2]
    assert [row["sold_qty"] for row in prepaid] == [min(500, expected), max(0, expected - 500)]
    assert sum(dec(row["reference_amount"]) for row in prepaid) == dec(expected) * dec("4.68")
    ledger = next(row for row in state["ledger"] if row["ledger_id"] == first["ledger_id"])
    alerts = [alert for alert in ledger["alerts"] if alert["text"].startswith("正常申报累计数量为")]
    assert alerts == (
        [
            {
                "level": "warning",
                "text": f"正常申报累计数量为 {sum(normal)} 件，请核对冲减记录；预退税可分配销量按实际净销量 {sold} 件封顶",
            }
        ]
        if sum(normal) < 0
        else []
    )


def test_optional_column_swaps_are_rejected_and_original_template_imports():
    original = Path(__file__).parents[1] / "docs/亚马逊9810核销台账模板-初版.xlsx"
    data, _ = excel.parse(original.read_bytes())
    assert len(data["exports"]) == 2
    book = load_workbook(BytesIO(excel.export(sample(), DEFAULT_SETTINGS)))
    spec = SCHEMA["receipts"]
    sheet = book[spec["sheet"]]
    columns = [
        next(f["column"] for f in spec["fields"] if f["key"] == key)
        for key in ["received_amount", "exchange_rate"]
    ]
    for row in range(spec["header_row"], sheet.max_row + 1):
        a, b = (sheet.cell(row, column) for column in columns)
        a.value, b.value = b.value, a.value
    output = BytesIO()
    book.save(output)
    with pytest.raises(RuleError, match="表头不匹配"):
        excel.parse(output.getvalue())


def test_high_precision_manual_amounts_export_as_text_and_remain_lossless():
    data = sample()
    data["exports"][0]["fob"] = "99999999999.12345678"
    output = excel.export(data, DEFAULT_SETTINGS)
    restored, _ = excel.parse(output)
    assert restored["exports"][0]["fob"] == data["exports"][0]["fob"]
    assert excel.merge(data, restored)[2] == 13
    book = load_workbook(BytesIO(output))
    spec = SCHEMA["exports"]
    column = next(f["column"] for f in spec["fields"] if f["key"] == "fob")
    assert book[spec["sheet"]].cell(spec["header_row"] + 1, column).data_type == "s"


def test_decimal_formatting_cannot_duplicate_records_after_export():
    data = sample()
    data["sales"][0]["gross"] = format(dec(data["sales"][0]["gross"]), ".2f")
    restored, _ = excel.parse(excel.export(data, DEFAULT_SETTINGS))
    _, added, skipped = excel.merge(data, restored)
    assert not added and skipped == 13
    assert clean("exports", {**data["exports"][0], "fob": "1E+9"})["fob"] == "1000000000"


@pytest.mark.parametrize("kind", ["taxes", "exceptions"])
@pytest.mark.parametrize("key", ["constructor", "__proto__", "toString"])
def test_reserved_names_are_not_existing_ledger_references(kind, key):
    data = sample()
    data[kind][0]["ledger_id"] = key
    with pytest.raises(RuleError, match="核销编号不存在"):
        validate_state(data, DEFAULT_SETTINGS)
